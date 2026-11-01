# Gera o modelo da minuta (campos entre chaves) a partir do .docx original da Vitalis.
# Uso: python3 gerar_modelo_minuta.py  (com minuta-original.docx na mesma pasta) -> modelo-contrato-vitalis.docx
# O modelo gerado sobe no painel em Contratos > Modelo da minuta. Não versionar o .docx: tem dados pessoais.
import docx, copy, re
frominuta-original.docx.oxml.ns import qn
d=docx.Document('minuta-original.docx')
T=d.tables

def set_par(p, text):
    runs=p.runs
    if runs:
        runs[0].text=text
        for r in runs[1:]: r._r.getparent().remove(r._r)
    else:
        p.add_run(text)
    # hyperlinks/other run containers
    for h in p._p.findall(qn('w:hyperlink')): p._p.remove(h)

def set_cell(cell, text):
    ps=cell.paragraphs
    set_par(ps[0], text)
    for p in ps[1:]: p._p.getparent().remove(p._p)

def cell(ti, ri, ci): return T[ti].rows[ri].cells[ci]

# 1. CONTRATANTE
def mescla(ti, ri, ini, fim):
    frominuta-original.docx.oxml import OxmlElement
    row=T[ti].rows[ri]
    tcs=[]
    for ci in range(ini, fim+1):
        tc=row.cells[ci]._tc
        if tc not in tcs: tcs.append(tc)
    first=tcs[0]; span=0; w=0
    for tc in tcs:
        g=tc.tcPr.find(qn('w:gridSpan')); span+=int(g.get(qn('w:val'))) if g is not None else 1
        tw=tc.tcPr.find(qn('w:tcW')); w+=int(tw.get(qn('w:w')))
    for tc in tcs[1:]: tc.getparent().remove(tc)
    g=first.tcPr.find(qn('w:gridSpan'))
    if g is None:
        g=OxmlElement('w:gridSpan'); first.tcPr.insert(1, g)
    g.set(qn('w:val'), str(span))
    first.tcPr.find(qn('w:tcW')).set(qn('w:w'), str(w))
    return docx.table._Cell(first, T[ti])
set_cell(mescla(0,1,1,2), '{razao_social}')
set_cell(cell(0,2,1), '{cnpj}')
set_cell(cell(0,3,1), '{representante_legal}')
set_cell(cell(0,4,1), '{endereco}')
set_cell(mescla(0,5,1,4), '{telefone}')
# e-mail oficial: acrescenta o e-mail logo após o rótulo
c=cell(0,6,0)
for p in c.paragraphs:
    if 'notificações:' in p.text and 'Todas' not in p.text.split('notificações:')[0]:
        r=p.add_run(' {email_notificacao}'); r.bold=True; break
else:
    p=c.paragraphs[0]; r=p.add_run(' {email_notificacao}'); r.bold=True

# 3. SERVIÇOS
servicos={ (0,16):'aso',(0,17):'pcmso',(0,19):'pgr',(0,20):'ltcat',(0,21):'psicossocial',(0,22):'ppp',(0,23):'insalubridade',
 (0,24):'periculosidade',(0,26):'esocial',(1,2):'cipa',(1,3):'brigada',(1,4):'primeiros_socorros',(1,5):'nrs',(1,6):'pericia',
 (1,8):'avulso',(1,9):'atualizacao_laudo'}
for (ti,ri),cod in servicos.items():
    set_cell(cell(ti,ri,0), '{x_%s}'%cod)
    set_cell(cell(ti,ri,3), '{v_%s}'%cod)
    set_cell(cell(ti,ri,4), '{c_%s}'%cod)
set_cell(cell(1,0,0), '{x_esocial_nao}')

# 4. PRAZO
set_cell(cell(1,12,1), '{x_p12} 12 (doze) meses')
set_cell(cell(1,12,2), '{x_p24} 24 (vinte e quatro) meses')
set_cell(cell(1,12,3), '{x_pout} Outro: {prazo_outro}')
set_cell(cell(1,13,1), '{inicio}\tTérmino previsto: {fim}')
# 6. PAGAMENTO
set_cell(cell(1,19,1), 'Dia {dia_vencimento} de cada mês')
set_cell(cell(1,20,1), '{x_boleto} Boleto bancário')
set_cell(cell(1,20,2), '{x_pix} PIX / Transferência bancária')
set_cell(cell(1,20,3), '{x_fout} Outro: {forma_outro}')

# Assinatura: local/data e nome do representante
feito={'data':False,'rep':False}
for p in d.paragraphs:
    t=p.text.strip()
    if not feito['data'] and re.match(r'^São Paulo, \d', t):
        set_par(p, 'São Paulo, {data_extenso}'); feito['data']=True
    elif not feito['rep'] and t=='Nome / Cargo':
        set_par(p, '{representante_legal}'); feito['rep']=True
print(feito)
d.save('modelo-contrato-vitalis.docx')
